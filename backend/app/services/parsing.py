"""Safe dataset parsing and inspection service for CSV, XLSX, and XLS workbooks."""
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
from typing import Any, BinaryIO, Optional
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import xlrd

from backend.app.config import settings
from backend.app.models.schemas import (
    DatasetInspectionResponse,
    ErrorCode,
    detect_canonical_columns,
)

# Supported workbook file extensions
SUPPORTED_EXTENSIONS: set[str] = {".csv", ".xlsx", ".xls"}

# OLE2 Compound File Binary Format magic bytes for legacy .xls workbooks
OLE2_MAGIC_BYTES: bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Regex pattern for identifying sensitive data columns to mask in preview
SENSITIVE_COLUMN_PATTERN: re.Pattern = re.compile(
    r"(password|secret|token|ssn|credit[-_]?card|cvv|api[-_]?key)",
    re.IGNORECASE,
)


class ParsingError(Exception):
    """Exception raised for dataset parsing and validation errors."""

    def __init__(self, code: ErrorCode | str, message: str, details: Optional[dict[str, Any]] = None):
        super().__init__(message)
        self.code = code if isinstance(code, ErrorCode) else ErrorCode(code)
        self.message = message
        self.details = details or {}


def sanitize_filename(raw_filename: Optional[str]) -> str:
    """Sanitize uploaded filename to prevent directory traversal and remove unsafe characters.

    Strips directory separators, null bytes, and non-alphanumeric characters (except '.', '_', '-').
    """
    if not raw_filename:
        return "uploaded_file"

    # Extract base filename to strip any path components
    name = Path(raw_filename).name
    name = name.replace("\x00", "")

    # Replace characters not matching safe pattern
    clean = re.sub(r"[^a-zA-Z0-9._-]", "_", name)
    clean = re.sub(r"_{2,}", "_", clean)
    clean = re.sub(r"\.{2,}", ".", clean)
    clean = clean.lstrip(".")

    return clean or "uploaded_file"


def get_file_extension(filename: str) -> str:
    """Extract lowercase file extension from filename."""
    return Path(filename).suffix.lower()


def validate_file_content(file_path: Path, ext: str) -> None:
    """Inspect file content and verify magic bytes / structure.

    Raises:
        ParsingError: If file is empty, corrupted, or does not match extension format.
    """
    if not file_path.exists():
        raise ParsingError(
            ErrorCode.DATASET_NOT_FOUND,
            f"File does not exist: {file_path}",
        )

    file_size = file_path.stat().st_size
    if file_size == 0:
        raise ParsingError(
            ErrorCode.EMPTY_FILE,
            "Uploaded file is empty (0 bytes)",
            {"file_size_bytes": 0},
        )

    ext_lower = ext.lower()
    if ext_lower not in SUPPORTED_EXTENSIONS:
        raise ParsingError(
            ErrorCode.UNSUPPORTED_FILE_TYPE,
            f"Unsupported file extension '{ext}'. Allowed extensions are: {sorted(SUPPORTED_EXTENSIONS)}",
            {"extension": ext, "supported_extensions": sorted(SUPPORTED_EXTENSIONS)},
        )

    if ext_lower == ".xlsx":
        if not zipfile.is_zipfile(file_path):
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                "Invalid XLSX file: file is not a valid ZIP/Office Open XML archive",
                {"extension": ext},
            )
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            wb.close()
        except Exception as err:
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                f"Failed to open XLSX workbook: {err}",
                {"extension": ext},
            ) from err

    elif ext_lower == ".xls":
        with open(file_path, "rb") as f:
            header = f.read(8)
        if header != OLE2_MAGIC_BYTES:
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                "Invalid XLS file: missing OLE2 compound document header",
                {"extension": ext},
            )
        try:
            wb = xlrd.open_workbook(file_path, on_demand=True)
            wb.release_resources()
        except Exception as err:
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                f"Failed to open XLS workbook: {err}",
                {"extension": ext},
            ) from err

    elif ext_lower == ".csv":
        # Check for binary content or null bytes in header sample
        with open(file_path, "rb") as f:
            sample = f.read(8192)
        if b"\x00" in sample:
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                "Invalid CSV file: binary data or null bytes detected",
                {"extension": ext},
            )
        # Attempt decoding with UTF-8 or Latin-1
        try:
            sample.decode("utf-8")
        except UnicodeDecodeError:
            try:
                sample.decode("latin-1")
            except Exception as err:
                raise ParsingError(
                    ErrorCode.UNSUPPORTED_FILE_TYPE,
                    f"Invalid CSV file encoding: {err}",
                    {"extension": ext},
                ) from err


def get_workbook_sheets(file_path: Path, ext: str) -> list[str]:
    """List sheet names available in the workbook.

    For CSV files, returns `["Sheet1"]`.
    """
    ext_lower = ext.lower()
    if ext_lower == ".csv":
        return ["Sheet1"]

    if ext_lower == ".xlsx":
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheets = list(wb.sheetnames)
        wb.close()
        return sheets

    if ext_lower == ".xls":
        wb = xlrd.open_workbook(file_path, on_demand=True)
        sheets = list(wb.sheet_names())
        wb.release_resources()
        return sheets

    return []


def read_sheet_dataframe(
    file_path: Path,
    ext: str,
    target_sheet: Optional[str] = None,
) -> tuple[pd.DataFrame, list[str]]:
    """Read target sheet into pandas DataFrame and return sheets list.

    Raises:
        ParsingError: If target sheet is not found or reading fails.
    """
    ext_lower = ext.lower()
    sheets = get_workbook_sheets(file_path, ext_lower)

    if ext_lower == ".csv":
        if target_sheet and target_sheet != "Sheet1":
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                f"Target sheet '{target_sheet}' not found. CSV files only have 'Sheet1'.",
                {"available_sheets": sheets, "requested_sheet": target_sheet},
            )
        try:
            df = pd.read_csv(file_path, encoding="utf-8")
        except UnicodeDecodeError:
            df = pd.read_csv(file_path, encoding="latin-1")
        except Exception as err:
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                f"Failed to parse CSV file: {err}",
            ) from err
        return df, sheets

    # For Excel workbooks
    if not sheets:
        raise ParsingError(
            ErrorCode.EMPTY_FILE,
            "Workbook contains no sheets",
        )

    if target_sheet:
        if target_sheet not in sheets:
            raise ParsingError(
                ErrorCode.UNSUPPORTED_FILE_TYPE,
                f"Target sheet '{target_sheet}' not found in workbook. Available sheets: {sheets}",
                {"available_sheets": sheets, "requested_sheet": target_sheet},
            )
        sheet_to_read = target_sheet
    else:
        sheet_to_read = sheets[0]

    try:
        if ext_lower == ".xlsx":
            df = pd.read_excel(file_path, sheet_name=sheet_to_read, engine="openpyxl")
        else:
            df = pd.read_excel(file_path, sheet_name=sheet_to_read, engine="xlrd")
    except Exception as err:
        raise ParsingError(
            ErrorCode.UNSUPPORTED_FILE_TYPE,
            f"Failed to read sheet '{sheet_to_read}': {err}",
            {"sheet": sheet_to_read},
        ) from err

    return df, sheets


def extract_available_months(df: pd.DataFrame, date_column: Optional[str]) -> list[str]:
    """Parse dates in the date column and identify sorted calendar months in YYYY-MM format."""
    if not date_column or date_column not in df.columns:
        return []

    try:
        parsed_dates = pd.to_datetime(df[date_column], errors="coerce")
        valid = parsed_dates.dropna()
        if valid.empty:
            return []
        months = sorted(list(valid.dt.strftime("%Y-%m").unique()))
        return months
    except Exception:
        return []


def extract_preview_rows(df: pd.DataFrame, max_rows: int = 10) -> list[dict[str, Any]]:
    """Convert first N rows to JSON-serializable dictionaries, masking sensitive fields."""
    preview_df = df.head(max_rows)
    preview: list[dict[str, Any]] = []

    for _, row in preview_df.iterrows():
        row_dict: dict[str, Any] = {}
        for col in df.columns:
            col_str = str(col)
            val = row[col]

            if SENSITIVE_COLUMN_PATTERN.search(col_str):
                row_dict[col_str] = "***"
            elif pd.isna(val):
                row_dict[col_str] = None
            elif isinstance(val, (pd.Timestamp, datetime, date)):
                row_dict[col_str] = str(val)
            elif hasattr(val, "item"):
                v = val.item()
                if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                    row_dict[col_str] = None
                else:
                    row_dict[col_str] = v
            elif isinstance(val, float) and (math.isnan(val) or math.isinf(val)):
                row_dict[col_str] = None
            else:
                row_dict[col_str] = val

        preview.append(row_dict)

    return preview


def inspect_uploaded_file(
    file_path: Path,
    filename: str,
    dataset_id: str,
    file_size_bytes: int,
    target_sheet: Optional[str] = None,
) -> DatasetInspectionResponse:
    """Inspect an uploaded workbook file and return inspection summary schema."""
    ext = get_file_extension(filename)
    validate_file_content(file_path, ext)

    df, sheets = read_sheet_dataframe(file_path, ext, target_sheet)

    # Detect canonical columns
    columns = [str(c).strip() for c in df.columns]
    detected_columns, missing_required_columns = detect_canonical_columns(columns)

    # Extract available months
    date_col = detected_columns.get("date")
    available_months = extract_available_months(df, date_col)

    # Count total rows and extract preview
    total_rows = int(len(df))
    preview_rows = extract_preview_rows(df, max_rows=10)

    return DatasetInspectionResponse(
        dataset_id=dataset_id,
        filename=filename,
        file_size_bytes=file_size_bytes,
        sheets=sheets,
        detected_columns=detected_columns,
        missing_required_columns=missing_required_columns,
        available_months=available_months,
        total_rows=total_rows,
        preview_rows=preview_rows,
    )


def get_dataset_dir(dataset_id: str, base_dir: Optional[Path | str] = None) -> Path:
    """Return the directory path for a staged dataset."""
    root = Path(base_dir or settings.upload_storage_dir)
    return root / dataset_id


def get_dataset_meta(dataset_id: str, base_dir: Optional[Path | str] = None) -> Optional[dict[str, Any]]:
    """Load metadata dictionary for a staged dataset if it exists."""
    meta_path = get_dataset_dir(dataset_id, base_dir) / "meta.json"
    if not meta_path.exists():
        return None
    try:
        with open(meta_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def save_uploaded_dataset(
    stream: BinaryIO,
    original_filename: Optional[str],
    dataset_id: str,
    max_size_bytes: int,
    base_dir: Optional[Path | str] = None,
) -> tuple[Path, str, int, str]:
    """Save an uploaded file stream into isolated staging storage.

    Returns:
        tuple[Path, str, int, str]: (raw_file_path, sanitized_filename, file_size_bytes, sha256_hash)

    Raises:
        ParsingError: If file exceeds size limit or is empty.
    """
    clean_filename = sanitize_filename(original_filename)
    ext = get_file_extension(clean_filename)

    if ext not in SUPPORTED_EXTENSIONS:
        raise ParsingError(
            ErrorCode.UNSUPPORTED_FILE_TYPE,
            f"Unsupported file extension '{ext}'. Allowed extensions are: {sorted(SUPPORTED_EXTENSIONS)}",
            {"extension": ext, "supported_extensions": sorted(SUPPORTED_EXTENSIONS)},
        )

    target_dir = get_dataset_dir(dataset_id, base_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    raw_file_path = target_dir / f"raw{ext}"

    # Invariant safety check: never overwrite demo.json
    demo_path = Path(__file__).resolve().parent.parent / "data" / "demo.json"
    if raw_file_path.resolve() == demo_path.resolve():
        raise RuntimeError("Safety invariant violation: attempted to overwrite demo.json")

    total_bytes = 0
    hasher = hashlib.sha256()

    try:
        with open(raw_file_path, "wb") as f_out:
            while True:
                chunk = stream.read(1024 * 1024)
                if not chunk:
                    break
                total_bytes += len(chunk)
                if total_bytes > max_size_bytes:
                    f_out.close()
                    shutil.rmtree(target_dir, ignore_errors=True)
                    raise ParsingError(
                        ErrorCode.FILE_TOO_LARGE,
                        f"File size exceeds maximum allowed size of {max_size_bytes} bytes",
                        {"max_upload_size_bytes": max_size_bytes, "uploaded_bytes": total_bytes},
                    )
                hasher.update(chunk)
                f_out.write(chunk)
    except Exception:
        if target_dir.exists():
            shutil.rmtree(target_dir, ignore_errors=True)
        raise

    if total_bytes == 0:
        shutil.rmtree(target_dir, ignore_errors=True)
        raise ParsingError(
            ErrorCode.EMPTY_FILE,
            "Uploaded file is empty (0 bytes)",
            {"filename": clean_filename},
        )

    sha256_hash = hasher.hexdigest()

    # Create metadata JSON in the directory
    meta = {
        "dataset_id": dataset_id,
        "filename": clean_filename,
        "original_filename": original_filename or clean_filename,
        "file_size_bytes": total_bytes,
        "sha256": sha256_hash,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "extension": ext,
        "raw_path": str(raw_file_path),
    }
    meta_path = target_dir / "meta.json"
    with open(meta_path, "w", encoding="utf-8") as f_meta:
        json.dump(meta, f_meta, indent=2)

    return raw_file_path, clean_filename, total_bytes, sha256_hash
