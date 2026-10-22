"""Tests for safe CSV, XLSX, and XLS workbook upload, validation, and inspection API."""
from datetime import datetime
import hashlib
import io
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
import openpyxl
import pandas as pd
import xlwt

from backend.app.config import settings
from backend.app.main import app
from backend.app.models.schemas import ErrorCode
from backend.app.services.parsing import get_dataset_dir, get_dataset_meta


REPO_ROOT = Path(__file__).resolve().parent.parent
DEMO_JSON_PATH = REPO_ROOT / "backend" / "app" / "data" / "demo.json"
DEMO_PROV_PATH = REPO_ROOT / "backend" / "app" / "data" / "demo-provenance.json"


@pytest.fixture
def client(tmp_path, monkeypatch):
    """Provide a TestClient with uploads redirected to an isolated temp directory."""
    monkeypatch.setattr(settings, "upload_storage_dir", str(tmp_path / "uploads"))
    return TestClient(app)


def compute_file_sha256(path: Path) -> str:
    """Compute sha256 checksum of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def create_sample_csv(rows: int = 5) -> bytes:
    """Generate sample CSV bytes with canonical columns across multiple calendar months."""
    lines = ["Invoice,StockCode,Description,Quantity,InvoiceDate"]
    for i in range(1, rows + 1):
        month = "01" if i % 2 == 1 else "02"
        lines.append(f"INV00{i},SKU00{i},Sample Item {i},{i * 2},2023-{month}-15 10:30:00")
    return "\n".join(lines).encode("utf-8")


def create_sample_xlsx(sheets: dict[str, list[list[str]]]) -> bytes:
    """Generate sample XLSX bytes with given sheet data."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    for sheet_name, data in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in data:
            ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def create_sample_xls(sheets: dict[str, list[list[str | int | float]]]) -> bytes:
    """Generate sample legacy XLS bytes using xlwt."""
    wb = xlwt.Workbook()
    for sheet_name, rows in sheets.items():
        ws = wb.add_sheet(sheet_name)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                ws.write(r_idx, c_idx, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# 1. Valid File Upload Tests (.csv, .xlsx, .xls)
# ==============================================================================

def test_upload_valid_csv(client, tmp_path):
    """Verify uploading a valid CSV returns canonical inspection schema, preview rows, and months."""
    csv_bytes = create_sample_csv(rows=4)
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("retail_orders.csv", csv_bytes, "text/csv")},
    )

    assert response.status_code == 200, response.text
    data = response.json()

    assert data["dataset_id"].startswith("ds_")
    assert data["filename"] == "retail_orders.csv"
    assert data["file_size_bytes"] == len(csv_bytes)
    assert data["sheets"] == ["Sheet1"]
    assert data["missing_required_columns"] == []

    # Verify detected columns mapping
    det = data["detected_columns"]
    assert det["invoice"] == "Invoice"
    assert det["sku"] == "StockCode"
    assert det["description"] == "Description"
    assert det["quantity"] == "Quantity"
    assert det["date"] == "InvoiceDate"

    # Verify available months extracted from InvoiceDate
    assert data["available_months"] == ["2023-01", "2023-02"]
    assert data["total_rows"] == 4
    assert len(data["preview_rows"]) == 4

    # Verify first preview row content
    first = data["preview_rows"][0]
    assert first["Invoice"] == "INV001"
    assert first["StockCode"] == "SKU001"
    assert first["Quantity"] == 2

    # Verify disk staging and metadata JSON
    dataset_dir = get_dataset_dir(data["dataset_id"], tmp_path / "uploads")
    assert dataset_dir.is_dir()
    assert (dataset_dir / "raw.csv").is_file()
    meta = get_dataset_meta(data["dataset_id"], tmp_path / "uploads")
    assert meta is not None
    assert meta["dataset_id"] == data["dataset_id"]
    assert meta["filename"] == "retail_orders.csv"
    assert meta["file_size_bytes"] == len(csv_bytes)
    assert meta["sha256"] == hashlib.sha256(csv_bytes).hexdigest()


def test_upload_valid_xlsx(client):
    """Verify uploading a multi-sheet XLSX returns sheet list and defaults to first sheet."""
    sheets_data = {
        "Year 2009-2010": [
            ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate"],
            ["489434", "85048", "15CM CHRISTMAS GLASS BALL", 12, "2009-12-01 07:45:00"],
        ],
        "Year 2010-2011": [
            ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate"],
            ["536365", "85123A", "WHITE HANGING HEART T-LIGHT HOLDER", 6, "2010-12-01 08:26:00"],
            ["536366", "22633", "HAND WARMER UNION JACK", 6, "2011-01-05 09:00:00"],
        ],
    }
    xlsx_bytes = create_sample_xlsx(sheets_data)

    response = client.post(
        "/api/datasets/upload",
        files={"file": ("online_retail.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200, response.text
    data = response.json()

    assert data["sheets"] == ["Year 2009-2010", "Year 2010-2011"]
    # Default sheet is Year 2009-2010 (1 row)
    assert data["total_rows"] == 1
    assert data["available_months"] == ["2009-12"]
    assert data["missing_required_columns"] == []
    assert data["preview_rows"][0]["Invoice"] in (489434, "489434")


def test_upload_valid_xlsx_with_target_sheet_param(client):
    """Verify target_sheet query parameter inspects specified sheet instead of first sheet."""
    sheets_data = {
        "Year 2009-2010": [
            ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate"],
            ["489434", "85048", "15CM CHRISTMAS GLASS BALL", 12, "2009-12-01 07:45:00"],
        ],
        "Year 2010-2011": [
            ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate"],
            ["536365", "85123A", "WHITE HANGING HEART T-LIGHT HOLDER", 6, "2010-12-01 08:26:00"],
            ["536366", "22633", "HAND WARMER UNION JACK", 6, "2011-06-15 09:00:00"],
        ],
    }
    xlsx_bytes = create_sample_xlsx(sheets_data)

    # Test via query parameter
    response = client.post(
        "/api/datasets/upload?target_sheet=Year+2010-2011",
        files={"file": ("online_retail.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["sheets"] == ["Year 2009-2010", "Year 2010-2011"]
    assert data["total_rows"] == 2
    assert data["available_months"] == ["2010-12", "2011-06"]

    # Test via multipart form data
    response_form = client.post(
        "/api/datasets/upload",
        files={"file": ("online_retail.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"target_sheet": "Year 2010-2011"},
    )
    assert response_form.status_code == 200, response_form.text
    assert response_form.json()["total_rows"] == 2


def test_upload_valid_xls(client):
    """Verify uploading legacy .xls workbook returns sheets, canonical columns, and preview."""
    sheets_data = {
        "Sheet1": [
            ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate"],
            ["541234", "71053", "WHITE METAL LANTERN", 6, "2011-03-01 12:00:00"],
        ]
    }
    xls_bytes = create_sample_xls(sheets_data)

    response = client.post(
        "/api/datasets/upload",
        files={"file": ("legacy_orders.xls", xls_bytes, "application/vnd.ms-excel")},
    )

    assert response.status_code == 200, response.text
    data = response.json()

    assert data["filename"] == "legacy_orders.xls"
    assert data["sheets"] == ["Sheet1"]
    assert data["total_rows"] == 1
    assert data["available_months"] == ["2011-03"]
    assert data["detected_columns"]["invoice"] == "Invoice"
    assert data["detected_columns"]["sku"] == "StockCode"


# ==============================================================================
# 2. Validation & Safety Rejection Tests
# ==============================================================================

@pytest.mark.parametrize("bad_name", ["test.txt", "payload.json", "script.py", "archive.zip", "doc.pdf"])
def test_upload_unsupported_extensions(client, bad_name):
    """Verify rejection of files without .csv, .xlsx, or .xls extensions."""
    response = client.post(
        "/api/datasets/upload",
        files={"file": (bad_name, b"some arbitrary content", "text/plain")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["code"] == ErrorCode.UNSUPPORTED_FILE_TYPE.value
    assert "Unsupported file extension" in detail["message"]


def test_upload_corrupted_xlsx(client):
    """Verify that a non-zip file masquerading as .xlsx is rejected."""
    fake_xlsx = b"Plain text masquerading as XLSX file"
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("fake.xlsx", fake_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["code"] == ErrorCode.UNSUPPORTED_FILE_TYPE.value
    assert "Invalid XLSX file" in detail["message"]


def test_upload_corrupted_xls(client):
    """Verify that a file masquerading as .xls without OLE2 header is rejected."""
    fake_xls = b"Plain text pretending to be XLS"
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("fake.xls", fake_xls, "application/vnd.ms-excel")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["code"] == ErrorCode.UNSUPPORTED_FILE_TYPE.value
    assert "Invalid XLS file" in detail["message"]


def test_upload_corrupted_binary_csv(client):
    """Verify that binary data with null bytes masquerading as CSV is rejected."""
    binary_csv = b"Invoice,StockCode\x00\x01\x02\x03\x04"
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("binary.csv", binary_csv, "text/csv")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["code"] == ErrorCode.UNSUPPORTED_FILE_TYPE.value
    assert "binary data" in detail["message"].lower() or "null bytes" in detail["message"].lower()


def test_upload_empty_file(client):
    """Verify that 0-byte file is rejected with EMPTY_FILE."""
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("empty.csv", b"", "text/csv")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["code"] == ErrorCode.EMPTY_FILE.value
    assert "empty (0 bytes)" in detail["message"]


def test_upload_oversized_file(client, monkeypatch):
    """Verify that file exceeding max_upload_size_bytes is rejected with HTTP 413 and FILE_TOO_LARGE."""
    # Set limit to 200 bytes
    monkeypatch.setattr(settings, "max_upload_size_bytes", 200)

    large_content = b"x" * 500
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("large.csv", large_content, "text/csv")},
    )

    assert response.status_code == 413
    detail = response.json()["detail"]
    assert detail["code"] == ErrorCode.FILE_TOO_LARGE.value
    assert "exceeds maximum allowed size" in detail["message"]


def test_upload_missing_required_columns(client):
    """Verify that an uploaded workbook missing required columns flags them in response."""
    # CSV missing 'quantity' and 'date'
    partial_csv = b"Invoice,StockCode,Description\nINV01,SKU01,A sample item\n"
    response = client.post(
        "/api/datasets/upload",
        files={"file": ("partial.csv", partial_csv, "text/csv")},
    )

    assert response.status_code == 200
    data = response.json()
    assert "quantity" in data["missing_required_columns"]
    assert "date" in data["missing_required_columns"]
    assert data["available_months"] == []
    assert data["detected_columns"]["invoice"] == "Invoice"
    assert data["detected_columns"]["sku"] == "StockCode"


def test_upload_nonexistent_target_sheet(client):
    """Verify that requesting a nonexistent sheet returns an informative error."""
    sheets_data = {
        "Year 2010": [
            ["Invoice", "StockCode", "Quantity", "InvoiceDate"],
            ["1001", "SKU1", 1, "2010-01-01"],
        ]
    }
    xlsx_bytes = create_sample_xlsx(sheets_data)

    response = client.post(
        "/api/datasets/upload?target_sheet=Year+3000",
        files={"file": ("orders.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code in (400, 422)
    detail = response.json()["detail"]
    assert "not found" in detail["message"].lower()


# ==============================================================================
# 3. Path Traversal & PII Sanitization Tests
# ==============================================================================

def test_upload_filename_sanitization(client, tmp_path):
    """Verify path traversal sequences and unsafe characters are stripped from filename."""
    csv_bytes = create_sample_csv(rows=1)
    traversal_name = "../../../../etc/passwd..test.csv"

    response = client.post(
        "/api/datasets/upload",
        files={"file": (traversal_name, csv_bytes, "text/csv")},
    )

    assert response.status_code == 200
    data = response.json()
    sanitized = data["filename"]

    # Traversal components must be removed
    assert "/" not in sanitized
    assert "\\" not in sanitized
    assert ".." not in sanitized
    assert sanitized.endswith(".csv")

    # Staged file must be safely in uploads directory
    dataset_dir = get_dataset_dir(data["dataset_id"], tmp_path / "uploads")
    assert (dataset_dir / "raw.csv").is_file()


def test_upload_pii_masking_in_preview(client):
    """Verify sensitive columns (e.g. password, secret, token, ssn) are masked in preview_rows."""
    csv_content = (
        "Invoice,StockCode,Quantity,InvoiceDate,password,secret_token,api_key\n"
        "INV001,SKU001,10,2023-01-01 10:00:00,secret123,tok_456,key_789\n"
    ).encode("utf-8")

    response = client.post(
        "/api/datasets/upload",
        files={"file": ("orders_with_pii.csv", csv_content, "text/csv")},
    )

    assert response.status_code == 200
    data = response.json()
    first_preview = data["preview_rows"][0]
    assert first_preview["password"] == "***"
    assert first_preview["secret_token"] == "***"
    assert first_preview["api_key"] == "***"
    assert first_preview["Invoice"] == "INV001"


# ==============================================================================
# 4. Invariant Protection: Never Overwrite demo.json
# ==============================================================================

def test_upload_does_not_alter_demo_json(client):
    """Verify that multiple upload calls under no circumstances modify backend/app/data/demo.json."""
    assert DEMO_JSON_PATH.exists(), f"demo.json not found at {DEMO_JSON_PATH}"
    initial_demo_hash = compute_file_sha256(DEMO_JSON_PATH)
    initial_demo_stat = DEMO_JSON_PATH.stat()

    if DEMO_PROV_PATH.exists():
        initial_prov_hash = compute_file_sha256(DEMO_PROV_PATH)
    else:
        initial_prov_hash = None

    # Perform multiple valid and invalid uploads
    client.post(
        "/api/datasets/upload",
        files={"file": ("test_run1.csv", create_sample_csv(rows=5), "text/csv")},
    )
    client.post(
        "/api/datasets/upload",
        files={"file": ("demo.json", b"malicious overwrite attempt", "application/json")},
    )
    client.post(
        "/api/datasets/upload",
        files={"file": ("demo.json.csv", create_sample_csv(rows=2), "text/csv")},
    )

    # Re-verify demo.json hash and size remain identical
    assert DEMO_JSON_PATH.exists()
    final_demo_hash = compute_file_sha256(DEMO_JSON_PATH)
    final_demo_stat = DEMO_JSON_PATH.stat()

    assert initial_demo_hash == final_demo_hash, "Invariant violated: demo.json was modified!"
    assert initial_demo_stat.st_size == final_demo_stat.st_size

    if initial_prov_hash:
        assert compute_file_sha256(DEMO_PROV_PATH) == initial_prov_hash
